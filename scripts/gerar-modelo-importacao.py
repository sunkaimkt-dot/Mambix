"""Gera public/modelos/modelo-importacao-mambix.xlsx (Combo 3).

Arquivo estatico de proposito: o modelo nao precisa de biblioteca nenhuma no
sistema. Rodar de novo so se mudar alguma coluna da importacao:

    pip install openpyxl
    python scripts/gerar-modelo-importacao.py

Os cabecalhos sao os mesmos que a tela reconhece sozinha (CAMPOS em
src/lib/importacao.ts).
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SAIDA = Path(__file__).resolve().parent.parent / "public" / "modelos" / "modelo-importacao-mambix.xlsx"
DATA = "DD/MM/YYYY"
VALOR = "#,##0.00"
CINZA = PatternFill("solid", fgColor="E2E8F0")

ABAS = {
    "Pagamentos": [
        ("Vencimento", DATA, 13), ("Código", "0", 9), ("Descrição", "@", 32), ("Valor", VALOR, 13),
        ("CFC", "0", 6), ("Competência", "@", 13), ("Data de pagamento", DATA, 18),
        ("Forma de pagamento", "@", 20), ("Banco", "@", 16),
    ],
    "Receitas": [
        ("Data", DATA, 13), ("Tipo de recebimento", "@", 22), ("Descrição", "@", 32), ("Valor", VALOR, 13), ("Banco", "@", 16),
    ],
    "Caixa Diário": [
        ("Data", DATA, 13), ("Tipo de venda", "@", 20), ("Valor", VALOR, 13),
    ],
}

INSTRUCOES = [
    ("Como usar", None),
    ("1. Preencha UMA aba por vez (Pagamentos, Receitas ou Caixa Diário). Não mude os títulos da linha 1.", None),
    ("2. No sistema: Lançamentos → Importar Excel / PDF. Escolha o tipo, a loja e o arquivo; a aba certa é escolhida sozinha.", None),
    ("3. Confira a pré-visualização: linhas com erro não entram; duplicadas vêm desmarcadas. Nada é gravado antes de você confirmar.", None),
    ("4. Errou? No histórico da tela, 'desfazer' apaga a importação inteira.", None),
    ("", None),
    ("Pagamentos", None),
    ("Vencimento *", "data (dd/mm/aaaa)."),
    ("Código *", "código de despesa de 1 a 100 (o mesmo da tela Pagamentos). Código sem nome em 'Códigos e listas' é recusado."),
    ("Descrição", "texto livre."),
    ("Valor *", "maior que zero. Aceita 1.234,56 ou R$ 1.234,56."),
    ("CFC", "1 fixa, 2 variável, 3 não operacional, 4 investimento. Em branco: o sistema sugere pelo grupo do código."),
    ("Competência", "mm/aaaa. Em branco: o mês do vencimento (DRE)."),
    ("Data de pagamento", "preenchida = conta já paga: entra com a baixa integral nessa data (DFC). Em branco = em aberto."),
    ("Forma de pagamento", "código ou nome, como em 'Códigos e listas'."),
    ("Banco", "nome igual ao cadastrado. Em branco: o banco escolhido na tela."),
    ("", None),
    ("Receitas (DFC)", None),
    ("Data *", "dia em que o dinheiro entrou."),
    ("Tipo de recebimento *", "código (1 a 20) ou nome, ex.: 3 ou CARTÃO CRÉDITO."),
    ("Descrição / Valor * / Banco", "como em Pagamentos."),
    ("", None),
    ("Caixa Diário (faturamento da DRE)", None),
    ("Data * / Tipo de venda * / Valor *", "o sistema guarda um valor por loja + dia + tipo; linhas iguais no arquivo são somadas. "
     "Se o dia já tiver valor, você escolhe pular ou substituir."),
    ("", None),
    ("* obrigatório", None),
]

wb = Workbook()
inst = wb.active
inst.title = "Instruções"
inst.column_dimensions["A"].width = 34
inst.column_dimensions["B"].width = 110
for i, (a, b) in enumerate(INSTRUCOES, start=1):
    inst.cell(i, 1, a)
    if b:
        inst.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    if b is None and a and not a[0].isdigit():
        inst.cell(i, 1).font = Font(bold=True, size=12 if a == "Como usar" else 11)

for nome, colunas in ABAS.items():
    ws = wb.create_sheet(nome)
    for j, (titulo, fmt, larg) in enumerate(colunas, start=1):
        c = ws.cell(1, j, titulo)
        c.font = Font(bold=True)
        c.fill = CINZA
        ws.column_dimensions[get_column_letter(j)].width = larg
        for i in range(2, 1001):
            ws.cell(i, j).number_format = fmt
    ws.freeze_panes = "A2"

SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print(f"ok: {SAIDA}")
